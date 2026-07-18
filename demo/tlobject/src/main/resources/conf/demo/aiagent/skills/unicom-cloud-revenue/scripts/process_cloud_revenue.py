#!/usr/bin/env python3
"""
联通云资源业务收入统计
按归属单位、云类型（资源池）统计各月收入，输出到单Sheet Excel文件。

用法:
    python process_cloud_revenue.py [选项]

选项:
    --dir <目录>        工作目录，默认为当前目录
    --months <范围>     月份范围，如 "1-5" 表示1月到5月，默认 "1-5"
    --year <年份>       年份，默认 2026
    --unit-file <文件>  单位顺序文件，默认 单位.xlsx
    --output <文件>     输出文件名，默认 云资源业务收入统计汇总_<年份>年<范围>月.xlsx
    --pattern <模板>    月度文件名模板，用 {year} 和 {month} 占位，
                        默认 "{year}年{month}月云资源业务明细表.xlsx"

文件结构要求:
    月度明细表: 至少8列 (A-H)
        A: 业务号码, B: 产品名称, C: 金额, D: 资源池(云类型),
        E: 客户名称, F: 客户经理, G: 归属业务单位, H: 收入(不含税)
    单位文件: 第1列为单位名称列表
"""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import os
import sys
import argparse
import re


def parse_args():
    parser = argparse.ArgumentParser(description='联通云资源业务收入统计')
    parser.add_argument('--dir', default='.', help='工作目录')
    parser.add_argument('--months', default='1-5', help='月份范围，如 1-5')
    parser.add_argument('--year', type=int, default=2026, help='年份')
    parser.add_argument('--unit-file', default='单位.xlsx', help='单位顺序文件')
    parser.add_argument('--output', default=None, help='输出文件名')
    parser.add_argument('--pattern', default='{year}年{month}月云资源业务明细表.xlsx',
                        help='月度文件名模板')
    return parser.parse_args()


def parse_month_range(range_str):
    """解析 '1-5' 为 [1,2,3,4,5]"""
    m = re.match(r'(\d+)-(\d+)', range_str)
    if m:
        return list(range(int(m.group(1)), int(m.group(2)) + 1))
    # 逗号分隔: "1,3,5"
    return [int(x.strip()) for x in range_str.split(',')]


def main():
    args = parse_args()
    BASE_DIR = args.dir
    YEAR = args.year
    MONTHS = parse_month_range(args.months)

    # 输出文件名
    if args.output:
        output_filename = args.output
    else:
        output_filename = f"云资源业务收入统计汇总_{YEAR}年{args.months}月.xlsx"

    TAX_RATE = 1.06

    # ===== 1. 读取单位顺序 =====
    unit_file_path = os.path.join(BASE_DIR, args.unit_file)
    if not os.path.exists(unit_file_path):
        print(f"错误: 找不到单位文件 {unit_file_path}")
        sys.exit(1)

    wb_units = openpyxl.load_workbook(unit_file_path)
    ws_units = wb_units.active
    unit_order = []
    for row in ws_units.iter_rows(min_row=1, max_row=ws_units.max_row, values_only=True):
        val = row[0]
        if val is not None and str(val).strip() and str(val).strip() != "单位":
            unit_order.append(str(val).strip())
    print(f"单位顺序 ({len(unit_order)}): {unit_order}")

    # ===== 2. 读取各月数据并统计 =====
    monthly_data = {}
    all_cloud_types = set()
    all_units_in_data = set()

    for month in MONTHS:
        fn = os.path.join(BASE_DIR, args.pattern.format(year=YEAR, month=month))
        if not os.path.exists(fn):
            print(f"警告: 找不到文件 {fn}，跳过")
            continue

        wb = openpyxl.load_workbook(fn, data_only=True)
        ws = wb.active
        unit_cloud_revenue = {}
        data_count = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            biz_num = row[0]
            if biz_num is None:
                break

            cloud_type = str(row[3]).strip() if row[3] else ""   # D 资源池 → 云类型
            unit = str(row[6]).strip() if row[6] else ""          # G 归属业务单位
            revenue_raw = row[7]                                   # H 收入(不含税)

            if revenue_raw is None:
                amount = row[2]
                if amount is not None:
                    revenue_raw = float(amount) / TAX_RATE

            if cloud_type and unit and revenue_raw is not None:
                revenue = float(revenue_raw)
                all_cloud_types.add(cloud_type)
                all_units_in_data.add(unit)

                if unit not in unit_cloud_revenue:
                    unit_cloud_revenue[unit] = {}
                unit_cloud_revenue[unit][cloud_type] = \
                    unit_cloud_revenue[unit].get(cloud_type, 0) + revenue

            data_count += 1

        monthly_data[month] = unit_cloud_revenue
        print(f"{month}月: {data_count} 条数据, {len(unit_cloud_revenue)} 个单位")

    # ===== 3. 确定单位列表 =====
    final_units = []
    for u in unit_order:
        if u in all_units_in_data:
            final_units.append(u)
    for u in sorted(all_units_in_data):
        if u not in final_units:
            final_units.append(u)

    cloud_type_list = sorted(all_cloud_types)
    cloud_type_list = [c for c in cloud_type_list if c]

    print(f"\n最终单位数: {len(final_units)}")
    print(f"云类型数: {len(cloud_type_list)}")
    print(f"云类型: {cloud_type_list}")

    # ===== 4. 汇总合计 =====
    total_unit_cloud = {}
    for month in MONTHS:
        if month not in monthly_data:
            continue
        for unit, cdata in monthly_data[month].items():
            if unit not in total_unit_cloud:
                total_unit_cloud[unit] = {}
            for ct, rev in cdata.items():
                total_unit_cloud[unit][ct] = total_unit_cloud[unit].get(ct, 0) + rev

    # ===== 5. 创建输出 =====
    output_wb = openpyxl.Workbook()
    ws = output_wb.active
    ws.title = f"{YEAR}年{args.months}月云资源收入统计"

    # 样式
    HEADER_MONTH_FONT = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    HEADER_MONTH_FILL = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    HEADER_SUB_FONT = Font(name='微软雅黑', size=9, bold=True, color='FFFFFF')
    HEADER_SUB_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
    DATA_FONT = Font(name='微软雅黑', size=9)
    DATA_ALIGN = Alignment(horizontal='right', vertical='center')
    UNIT_ALIGN = Alignment(horizontal='center', vertical='center')
    TOTAL_FONT = Font(name='微软雅黑', size=9, bold=True)
    TOTAL_FILL = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
    THIN_BORDER = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    NUM_FMT = '#,##0.00'

    active_months = [m for m in MONTHS if m in monthly_data]
    NUM_CLOUD = len(cloud_type_list)
    COLS_PER_MONTH = NUM_CLOUD + 1
    TOTAL_COLS = 1 + len(active_months) * COLS_PER_MONTH + COLS_PER_MONTH

    # 标题行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
    title_cell = ws.cell(row=1, column=1,
                         value=f"{YEAR}年{args.months}月 云资源业务收入统计（按归属单位×云类型）")
    title_cell.font = Font(name='微软雅黑', size=14, bold=True, color='1F4E79')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # 表头第一行: 归属单位 + 月份合并 + 合计
    ws.row_dimensions[2].height = 22
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
    cell = ws.cell(row=2, column=1, value='归属单位')
    cell.font = HEADER_MONTH_FONT
    cell.fill = HEADER_MONTH_FILL
    cell.alignment = HEADER_ALIGN
    cell.border = THIN_BORDER
    ws.cell(row=3, column=1).font = HEADER_MONTH_FONT
    ws.cell(row=3, column=1).fill = HEADER_MONTH_FILL
    ws.cell(row=3, column=1).border = THIN_BORDER

    for mi, month in enumerate(active_months):
        start_col = 2 + mi * COLS_PER_MONTH
        end_col = start_col + COLS_PER_MONTH - 1
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
        cell = ws.cell(row=2, column=start_col, value=f'{month}月')
        cell.font = HEADER_MONTH_FONT
        cell.fill = HEADER_MONTH_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
        for c in range(start_col, end_col + 1):
            ws.cell(row=2, column=c).border = THIN_BORDER
            ws.cell(row=2, column=c).font = HEADER_MONTH_FONT
            ws.cell(row=2, column=c).fill = HEADER_MONTH_FILL

    total_start_col = 2 + len(active_months) * COLS_PER_MONTH
    total_end_col = total_start_col + COLS_PER_MONTH - 1
    ws.merge_cells(start_row=2, start_column=total_start_col, end_row=2, end_column=total_end_col)
    cell = ws.cell(row=2, column=total_start_col, value=f'合计（{args.months}月）')
    cell.font = HEADER_MONTH_FONT
    cell.fill = HEADER_MONTH_FILL
    cell.alignment = HEADER_ALIGN
    cell.border = THIN_BORDER
    for c in range(total_start_col, total_end_col + 1):
        ws.cell(row=2, column=c).border = THIN_BORDER
        ws.cell(row=2, column=c).font = HEADER_MONTH_FONT
        ws.cell(row=2, column=c).fill = HEADER_MONTH_FILL

    # 表头第二行: 云类型子标题
    ws.row_dimensions[3].height = 20

    def write_sub_headers(start_col, sub_label='小计'):
        for j, ct in enumerate(cloud_type_list):
            cell = ws.cell(row=3, column=start_col + j, value=ct)
            cell.font = HEADER_SUB_FONT
            cell.fill = HEADER_SUB_FILL
            cell.alignment = HEADER_ALIGN
            cell.border = THIN_BORDER
        cell_sub = ws.cell(row=3, column=start_col + NUM_CLOUD, value=sub_label)
        cell_sub.font = HEADER_SUB_FONT
        cell_sub.fill = HEADER_SUB_FILL
        cell_sub.alignment = HEADER_ALIGN
        cell_sub.border = THIN_BORDER

    for mi, month in enumerate(active_months):
        write_sub_headers(2 + mi * COLS_PER_MONTH, '小计')
    write_sub_headers(total_start_col, '总计')

    # 数据行
    DATA_START_ROW = 4
    grand_col_totals = {}

    for i, unit in enumerate(final_units):
        row_num = DATA_START_ROW + i
        ws.row_dimensions[row_num].height = 18

        cell_unit = ws.cell(row=row_num, column=1, value=unit)
        cell_unit.font = DATA_FONT
        cell_unit.alignment = UNIT_ALIGN
        cell_unit.border = THIN_BORDER

        for mi, month in enumerate(active_months):
            start_col = 2 + mi * COLS_PER_MONTH
            udata = monthly_data[month].get(unit, {})
            row_subtotal = 0.0
            for j, ct in enumerate(cloud_type_list):
                val = udata.get(ct, 0)
                row_subtotal += val
                cell_val = ws.cell(row=row_num, column=start_col + j, value=round(val, 2))
                cell_val.font = DATA_FONT
                cell_val.alignment = DATA_ALIGN
                cell_val.number_format = NUM_FMT
                cell_val.border = THIN_BORDER

            cell_sub = ws.cell(row=row_num, column=start_col + NUM_CLOUD, value=round(row_subtotal, 2))
            cell_sub.font = Font(name='微软雅黑', size=9, bold=True)
            cell_sub.alignment = DATA_ALIGN
            cell_sub.number_format = NUM_FMT
            cell_sub.border = THIN_BORDER

        # 合计区
        tdata = total_unit_cloud.get(unit, {})
        row_grand_total = 0.0
        for j, ct in enumerate(cloud_type_list):
            val = tdata.get(ct, 0)
            row_grand_total += val
            grand_col_totals[ct] = grand_col_totals.get(ct, 0) + val
            cell_val = ws.cell(row=row_num, column=total_start_col + j, value=round(val, 2))
            cell_val.font = DATA_FONT
            cell_val.alignment = DATA_ALIGN
            cell_val.number_format = NUM_FMT
            cell_val.border = THIN_BORDER

        grand_col_totals['_grand'] = grand_col_totals.get('_grand', 0) + row_grand_total
        cell_gt = ws.cell(row=row_num, column=total_start_col + NUM_CLOUD, value=round(row_grand_total, 2))
        cell_gt.font = Font(name='微软雅黑', size=9, bold=True)
        cell_gt.alignment = DATA_ALIGN
        cell_gt.number_format = NUM_FMT
        cell_gt.border = THIN_BORDER

    # 合计行
    total_row = DATA_START_ROW + len(final_units)
    ws.row_dimensions[total_row].height = 20

    cell_tl = ws.cell(row=total_row, column=1, value='合计')
    cell_tl.font = TOTAL_FONT
    cell_tl.fill = TOTAL_FILL
    cell_tl.alignment = UNIT_ALIGN
    cell_tl.border = THIN_BORDER

    for mi, month in enumerate(active_months):
        start_col = 2 + mi * COLS_PER_MONTH
        month_subtotal = 0.0
        for j, ct in enumerate(cloud_type_list):
            col_sum = sum(monthly_data[month].get(u, {}).get(ct, 0) for u in final_units)
            month_subtotal += col_sum
            cell_val = ws.cell(row=total_row, column=start_col + j, value=round(col_sum, 2))
            cell_val.font = TOTAL_FONT
            cell_val.fill = TOTAL_FILL
            cell_val.alignment = DATA_ALIGN
            cell_val.number_format = NUM_FMT
            cell_val.border = THIN_BORDER

        cell_sub = ws.cell(row=total_row, column=start_col + NUM_CLOUD, value=round(month_subtotal, 2))
        cell_sub.font = TOTAL_FONT
        cell_sub.fill = TOTAL_FILL
        cell_sub.alignment = DATA_ALIGN
        cell_sub.number_format = NUM_FMT
        cell_sub.border = THIN_BORDER

    for j, ct in enumerate(cloud_type_list):
        val = grand_col_totals.get(ct, 0)
        cell_val = ws.cell(row=total_row, column=total_start_col + j, value=round(val, 2))
        cell_val.font = TOTAL_FONT
        cell_val.fill = TOTAL_FILL
        cell_val.alignment = DATA_ALIGN
        cell_val.number_format = NUM_FMT
        cell_val.border = THIN_BORDER

    cell_gt = ws.cell(row=total_row, column=total_start_col + NUM_CLOUD,
                      value=round(grand_col_totals.get('_grand', 0), 2))
    cell_gt.font = TOTAL_FONT
    cell_gt.fill = TOTAL_FILL
    cell_gt.alignment = DATA_ALIGN
    cell_gt.number_format = NUM_FMT
    cell_gt.border = THIN_BORDER

    # 列宽与冻结
    ws.column_dimensions['A'].width = 14
    for c in range(2, TOTAL_COLS + 1):
        ws.column_dimensions[get_column_letter(c)].width = 15
    ws.freeze_panes = 'B4'

    # 保存
    output_path = os.path.join(BASE_DIR, output_filename)
    output_wb.save(output_path)
    total_revenue = grand_col_totals.get('_grand', 0)
    print(f"\n输出文件: {output_path}")
    print(f"总合计收入: {total_revenue:,.2f}")
    return output_path


if __name__ == '__main__':
    main()
